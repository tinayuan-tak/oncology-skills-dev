#!/usr/bin/env python3
"""Generate PowerPoint snapshots from target-profile figure outputs.

Usage:
    python generate_snapshot.py \
        --input /path/to/KRAS-COADREAD/2026-09-10__2.0.0__23ad0c5 \
        --topics dependency,genomic_alteration \
        --template /path/to/template.potx \
        --output KRAS-COADREAD-snapshot.pptx
"""

from __future__ import annotations

import io
import re
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Optional

import click
import yaml

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.dml.color import RGBColor
except ImportError:
    raise ImportError("python-pptx is required. Install with: pip install python-pptx")

try:
    import cairosvg
    HAS_CAIROSVG = True
except ImportError:
    HAS_CAIROSVG = False

try:
    from PIL import Image
    HAS_PIL = True
except ImportError:
    HAS_PIL = False


SKILL_DIR = Path(__file__).parent.parent
DEFAULT_CONFIG = SKILL_DIR / "topics.yaml"


@dataclass
class FigureSpec:
    """A figure to include on a slide."""
    path: str
    caption: str


@dataclass
class SlideSpec:
    """A slide specification with title, subtitle, and figures."""
    title: str
    subtitle: str
    figures: list[FigureSpec]


@dataclass
class TopicSpec:
    """A topic containing multiple slides."""
    key: str
    display_name: str
    description: str
    hero: str
    slides: list[SlideSpec]


def load_topics_config(config_path: Path) -> dict[str, TopicSpec]:
    """Load topic configuration from YAML."""
    with config_path.open() as f:
        data = yaml.safe_load(f)

    topics = {}
    for key, cfg in data.get("topics", {}).items():
        slides = []
        for slide_cfg in cfg.get("slides", []):
            figures = [
                FigureSpec(path=fig["path"], caption=fig.get("caption", ""))
                for fig in slide_cfg.get("figures", [])
            ]
            slides.append(SlideSpec(
                title=slide_cfg.get("title", ""),
                subtitle=slide_cfg.get("subtitle", ""),
                figures=figures,
            ))
        topics[key] = TopicSpec(
            key=key,
            display_name=cfg.get("display_name", key),
            description=cfg.get("description", ""),
            hero=cfg.get("hero", ""),
            slides=slides,
        )
    return topics


def parse_run_dir(run_dir: Path) -> tuple[str, str]:
    """Extract gene and indication from run directory path.

    Expected format: .../GENE-INDICATION/date__version__hash/
    The GENE-INDICATION folder should be the parent of the run folder (which has date__version__hash format).
    """
    run_folder_pattern = re.compile(r"^\d{4}-\d{2}-\d{2}__[\d.]+__[a-f0-9]+$")

    if run_folder_pattern.match(run_dir.name):
        parent = run_dir.parent.name
        match = re.match(r"^([A-Z0-9]+)-([A-Z]+)$", parent, re.IGNORECASE)
        if match:
            return match.group(1).upper(), match.group(2).upper()

    for part in reversed(run_dir.parts):
        if part.startswith("OneDrive") or part.startswith("."):
            continue
        match = re.match(r"^([A-Z0-9]+)-([A-Z]+)$", part, re.IGNORECASE)
        if match and len(match.group(1)) <= 10 and len(match.group(2)) <= 10:
            return match.group(1).upper(), match.group(2).upper()

    raise ValueError(f"Could not parse gene-indication from path: {run_dir}")


def svg_to_png(svg_path: Path, scale: float = 2.0) -> bytes:
    """Convert SVG to PNG bytes using cairosvg."""
    if not HAS_CAIROSVG:
        raise ImportError("cairosvg is required for SVG conversion. Install with: pip install cairosvg")
    return cairosvg.svg2png(url=str(svg_path), scale=scale)


def get_image_bytes(figure_path: Path) -> tuple[bytes, str]:
    """Get image bytes and format, converting SVG if needed."""
    suffix = figure_path.suffix.lower()

    if suffix == ".svg":
        return svg_to_png(figure_path), "png"
    elif suffix == ".png":
        return figure_path.read_bytes(), "png"
    elif suffix in (".jpg", ".jpeg"):
        return figure_path.read_bytes(), "jpeg"
    else:
        raise ValueError(f"Unsupported image format: {suffix}")


def substitute_placeholders(text: str, gene: str, indication: str) -> str:
    """Replace placeholders in text with actual values."""
    return text.replace("{gene}", gene).replace("{indication}", indication)


def add_title_slide(prs: Presentation, gene: str, indication: str, topics: list[str]) -> None:
    """Add a title slide."""
    slide_layout = prs.slide_layouts[1]  # Title slide layout
    slide = prs.slides.add_slide(slide_layout)

    title = slide.shapes.title
    if title:
        title.text = f"{gene} in {indication}"

    try:
        subtitle = slide.placeholders[1]
        subtitle.text = f"Target Profile Snapshot: {', '.join(topics)}"
    except KeyError:
        for shape in slide.placeholders:
            if shape.placeholder_format.idx != 0:
                shape.text = f"Target Profile Snapshot: {', '.join(topics)}"
                break


def add_content_slide(
    prs: Presentation,
    title: str,
    subtitle: str,
    figures: list[tuple[Path, str]],
    gene: str,
    indication: str,
) -> None:
    """Add a content slide with figures.

    Args:
        prs: Presentation object
        title: Slide title (with placeholders)
        subtitle: Slide subtitle (with placeholders)
        figures: List of (figure_path, caption) tuples
        gene: Gene symbol for placeholder substitution
        indication: Indication code for placeholder substitution
    """
    slide_layout = prs.slide_layouts[8]  # Standard 1-Column Text layout
    slide = prs.slides.add_slide(slide_layout)

    slide_width = prs.slide_width
    slide_height = prs.slide_height

    title_text = substitute_placeholders(title, gene, indication)
    subtitle_text = substitute_placeholders(subtitle, gene, indication)

    title_box = slide.shapes.add_textbox(
        Inches(0.5), Inches(0.3), slide_width - Inches(1), Inches(0.6)
    )
    title_frame = title_box.text_frame
    title_para = title_frame.paragraphs[0]
    title_para.text = title_text
    title_para.font.size = Pt(28)
    title_para.font.bold = True
    title_para.font.color.rgb = RGBColor(0, 51, 102)  # Dark blue

    if subtitle_text:
        subtitle_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(0.85), slide_width - Inches(1), Inches(0.4)
        )
        subtitle_frame = subtitle_box.text_frame
        subtitle_para = subtitle_frame.paragraphs[0]
        subtitle_para.text = subtitle_text
        subtitle_para.font.size = Pt(16)
        subtitle_para.font.color.rgb = RGBColor(102, 102, 102)  # Gray

    if not figures:
        return

    content_top = Inches(1.5)
    content_height = slide_height - content_top - Inches(0.5)
    content_width = slide_width - Inches(1)

    n_figures = len(figures)
    if n_figures == 1:
        fig_width = content_width * 0.8
        fig_height = content_height * 0.85
        positions = [(Inches(0.5) + (content_width - fig_width) / 2, content_top)]
    elif n_figures == 2:
        fig_width = (content_width - Inches(0.5)) / 2
        fig_height = content_height * 0.85
        positions = [
            (Inches(0.5), content_top),
            (Inches(0.5) + fig_width + Inches(0.5), content_top),
        ]
    else:
        cols = 2
        rows = (n_figures + 1) // 2
        fig_width = (content_width - Inches(0.5)) / 2
        fig_height = (content_height - Inches(0.3) * (rows - 1)) / rows * 0.85
        positions = []
        for i in range(n_figures):
            row, col = divmod(i, cols)
            x = Inches(0.5) + col * (fig_width + Inches(0.5))
            y = content_top + row * (fig_height + Inches(0.5))
            positions.append((x, y))

    for i, (fig_path, caption) in enumerate(figures):
        if i >= len(positions):
            break
        x, y = positions[i]

        try:
            img_bytes, img_format = get_image_bytes(fig_path)
            img_stream = io.BytesIO(img_bytes)

            if HAS_PIL:
                with Image.open(io.BytesIO(img_bytes)) as img:
                    img_width, img_height = img.size
                aspect = img_width / img_height
            else:
                aspect = 1.5

            if fig_width / fig_height > aspect:
                actual_height = fig_height
                actual_width = fig_height * aspect
            else:
                actual_width = fig_width
                actual_height = fig_width / aspect

            slide.shapes.add_picture(img_stream, x, y, width=actual_width, height=actual_height)

            if caption:
                caption_box = slide.shapes.add_textbox(
                    x, y + actual_height + Inches(0.05), actual_width, Inches(0.3)
                )
                caption_frame = caption_box.text_frame
                caption_para = caption_frame.paragraphs[0]
                caption_para.text = caption
                caption_para.font.size = Pt(10)
                caption_para.font.color.rgb = RGBColor(102, 102, 102)
                caption_para.alignment = PP_ALIGN.CENTER

        except Exception as e:
            click.echo(f"  Warning: Could not add figure {fig_path}: {e}", err=True)


def load_presentation_from_template(template_path: Path) -> Presentation:
    """Load a presentation from template, handling both .pptx and .potx files."""
    import shutil
    import zipfile

    suffix = template_path.suffix.lower()

    if suffix == ".potx":
        temp_dir = tempfile.mkdtemp()
        temp_pptx = Path(temp_dir) / "presentation.pptx"

        with zipfile.ZipFile(template_path, 'r') as zf_in:
            with zipfile.ZipFile(temp_pptx, 'w', zipfile.ZIP_DEFLATED) as zf_out:
                for item in zf_in.infolist():
                    data = zf_in.read(item.filename)
                    if item.filename == '[Content_Types].xml':
                        data = data.decode('utf-8').replace(
                            'application/vnd.openxmlformats-officedocument.presentationml.template.main+xml',
                            'application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml'
                        ).encode('utf-8')
                    zf_out.writestr(item, data)

        prs = Presentation(str(temp_pptx))
        shutil.rmtree(temp_dir, ignore_errors=True)
        return prs

    return Presentation(str(template_path))


def generate_snapshot(
    run_dir: Path,
    topics: list[str],
    template_path: Path,
    output_path: Path,
    config_path: Path,
) -> None:
    """Generate a PowerPoint snapshot from target-profile outputs."""
    gene, indication = parse_run_dir(run_dir)
    click.echo(f"Generating snapshot for {gene} in {indication}")

    figures_dir = run_dir / "figures"
    if not figures_dir.exists():
        raise FileNotFoundError(f"Figures directory not found: {figures_dir}")

    topic_configs = load_topics_config(config_path)

    if "all" in topics:
        selected_topics = list(topic_configs.keys())
    else:
        selected_topics = []
        for t in topics:
            if t in topic_configs:
                selected_topics.append(t)
            else:
                click.echo(f"  Warning: Unknown topic '{t}', skipping", err=True)

    if not selected_topics:
        raise ValueError("No valid topics selected")

    click.echo(f"  Topics: {', '.join(selected_topics)}")

    prs = load_presentation_from_template(template_path)

    add_title_slide(prs, gene, indication, selected_topics)

    for topic_key in selected_topics:
        topic = topic_configs[topic_key]
        click.echo(f"  Processing topic: {topic.display_name}")

        for slide_spec in topic.slides:
            figures_to_add = []
            for fig_spec in slide_spec.figures:
                fig_path = figures_dir / fig_spec.path
                if fig_path.exists():
                    figures_to_add.append((fig_path, fig_spec.caption))
                else:
                    png_path = fig_path.with_suffix(".png")
                    if png_path.exists():
                        figures_to_add.append((png_path, fig_spec.caption))
                    else:
                        click.echo(f"    Warning: Figure not found: {fig_spec.path}", err=True)

            if figures_to_add:
                add_content_slide(
                    prs,
                    slide_spec.title,
                    slide_spec.subtitle,
                    figures_to_add,
                    gene,
                    indication,
                )

    prs.save(str(output_path))
    click.echo(f"  Saved: {output_path}")


@click.command()
@click.option(
    "--input", "-i", "input_dir",
    type=click.Path(exists=True, file_okay=False, path_type=Path),
    required=True,
    help="Path to target-profile run directory (e.g., .../KRAS-COADREAD/2026-09-10__2.0.0__hash)",
)
@click.option(
    "--topics", "-t",
    default="all",
    help="Comma-separated topics to include, or 'all' (default: all)",
)
@click.option(
    "--template",
    type=click.Path(exists=True, dir_okay=False, path_type=Path),
    required=True,
    help="Path to PowerPoint template (.potx)",
)
@click.option(
    "--output", "-o",
    type=click.Path(dir_okay=False, path_type=Path),
    default=None,
    help="Output PPTX path (default: {gene}-{indication}-snapshot.pptx)",
)
@click.option(
    "--config", "-c",
    type=click.Path(exists=True, dir_okay=False, path_type=Path),
    default=None,
    help="Topic configuration YAML (default: topics.yaml in skill dir)",
)
def main(
    input_dir: Path,
    topics: str,
    template: Path,
    output: Optional[Path],
    config: Optional[Path],
) -> None:
    """Generate PowerPoint snapshot from target-profile outputs."""
    topic_list = [t.strip() for t in topics.split(",")]

    config_path = config or DEFAULT_CONFIG

    if output is None:
        try:
            gene, indication = parse_run_dir(input_dir)
            output = Path(f"{gene}-{indication}-snapshot.pptx")
        except ValueError:
            output = Path("snapshot.pptx")

    generate_snapshot(
        run_dir=input_dir,
        topics=topic_list,
        template_path=template,
        output_path=output,
        config_path=config_path,
    )


if __name__ == "__main__":
    main()
